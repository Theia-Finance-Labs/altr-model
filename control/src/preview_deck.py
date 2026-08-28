"""Render the generated .pptx to per-slide PNGs for visual QA — no GUI / LibreOffice needed.

Reads the ACTUAL shapes (exact EMU positions, text, embedded formula images, tables) from the
pptx via python-pptx and paints them onto a 13.333x7.5in matplotlib canvas. Text wrapping is
estimated from font size and box width, so vertical overflow and text/formula overlap are visible.
Light box outlines reveal geometry.

Run:  .venv/bin/python control/src/preview_deck.py
Out:  control/figures/preview/slideNN.png
"""
import io
import os
import textwrap

import matplotlib

matplotlib.use("Agg")
matplotlib.rcParams["text.parse_math"] = False  # body text has literal "$1.1T"; formula PNGs are pre-rendered
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "..", "ALTR_Asset-Level_Transition_Risk_Methodology_v1.pptx")
OUT = os.path.join(HERE, "..", "figures", "preview")
os.makedirs(OUT, exist_ok=True)

SW_IN, SH_IN = 13.333, 7.5
EMU = 914400.0


def emu_in(v):
    return (v or 0) / EMU


def para_text(p):
    return "".join(r.text for r in p.runs)


def para_size(p, default=13.5):
    sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
    return max(sizes) if sizes else default


def draw_textframe(ax, shape):
    left, top = emu_in(shape.left), emu_in(shape.top)
    width = emu_in(shape.width)
    # light outline of the text box
    ax.add_patch(Rectangle((left, SH_IN - top), width, -emu_in(shape.height),
                           fill=False, ec="#cccccc", lw=0.4))
    y = top
    for p in shape.text_frame.paragraphs:
        txt = para_text(p)
        if not txt:
            y += 0.10
            continue
        size = para_size(p)
        # estimate chars-per-line from box width and font size
        cpl = max(8, int(width * 142 / size))
        line_h = 1.28 * size / 72.0
        for line in textwrap.wrap(txt, width=cpl) or [""]:
            ax.text(left + 0.02, SH_IN - y, line, fontsize=size, va="top", ha="left",
                    family="sans-serif", color="black")
            y += line_h
        y += max(0.04, (p.space_after.pt / 72.0) if p.space_after else 0.05)


def draw_picture(ax, shape):
    left, top = emu_in(shape.left), emu_in(shape.top)
    width, height = emu_in(shape.width), emu_in(shape.height)
    try:
        img = plt.imread(io.BytesIO(shape.image.blob))
        ax.imshow(img, extent=(left, left + width, SH_IN - top - height, SH_IN - top),
                  aspect="auto", zorder=5)
    except Exception:
        ax.add_patch(Rectangle((left, SH_IN - top - height), width, height,
                               fill=False, ec="red", lw=0.8))


def draw_table(ax, shape):
    tbl = shape.table
    left, top = emu_in(shape.left), emu_in(shape.top)
    col_w = [emu_in(c.width) for c in tbl.columns]
    row_h = [emu_in(r.height) for r in tbl.rows]
    y = top
    for i, r in enumerate(tbl.rows):
        x = left
        for j, c in enumerate(r.cells):
            ax.add_patch(Rectangle((x, SH_IN - y - row_h[i]), col_w[j], row_h[i],
                                   fill=False, ec="#999999", lw=0.4))
            size = 11
            cpl = max(8, int(col_w[j] * 142 / size))
            ty = y + 0.06
            for line in textwrap.wrap(c.text, width=cpl) or [""]:
                ax.text(x + 0.05, SH_IN - ty, line, fontsize=size, va="top", ha="left",
                        family="sans-serif", color="black", fontweight="bold" if i == 0 else "normal")
                ty += 1.25 * size / 72.0
            x += col_w[j]
        y += row_h[i]


def main():
    prs = Presentation(PPTX)
    for idx, slide in enumerate(prs.slides, 1):
        fig = plt.figure(figsize=(SW_IN, SH_IN), dpi=110)
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, SW_IN)
        ax.set_ylim(0, SH_IN)
        ax.axis("off")
        ax.add_patch(Rectangle((0, 0), SW_IN, SH_IN, fill=True, fc="white", ec="black", lw=1.0, zorder=0))
        for shape in slide.shapes:
            if shape.shape_type == 13 or shape.shape_type == 18 or getattr(shape, "image", None) is not None and shape.has_text_frame is False:
                pass
            if shape.has_text_frame and shape.text_frame.text.strip():
                draw_textframe(ax, shape)
            elif getattr(shape, "has_table", False) and shape.has_table:
                draw_table(ax, shape)
            else:
                # picture or connector
                try:
                    if shape.image is not None:
                        draw_picture(ax, shape)
                except Exception:
                    pass
        fig.savefig(os.path.join(OUT, f"slide{idx:02d}.png"), dpi=110)
        plt.close(fig)
    print(f"Rendered {idx} slide previews -> {OUT}")


if __name__ == "__main__":
    main()
